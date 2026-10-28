import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, List, Any, Optional

class ISOExcelService:
    @staticmethod
    def load_master_checklist(file_path: Optional[str] = None) -> List[Dict[str, Any]]:
        """Loads ISO controls from the master workbook or custom uploaded file."""
        if not file_path or not os.path.exists(file_path):
            file_path = os.path.join("data", "iso_checklists", "ISO_27002_2022_Master_Checklist.xlsx")
            
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1]]
        header_map = {}
        for idx, h in enumerate(headers):
            if not h:
                continue
            h_clean = str(h).lower().strip().replace(" ", "_").replace("-", "_")
            header_map[h_clean] = idx

        # Standard column fallback resolutions
        def get_val(row, *aliases):
            for alias in aliases:
                for k, idx in header_map.items():
                    if alias in k:
                        val = row[idx].value
                        return str(val).strip() if val is not None else ""
            return ""

        controls = []
        for row in ws.iter_rows(min_row=2):
            if not any(cell.value for cell in row):
                continue
            
            c_id = get_val(row, "control_id", "id", "nummer", "control")
            c_title = get_val(row, "title", "bezeichnung", "name")
            c_req = get_val(row, "plain_english", "requirement", "anforderung", "description")
            c_chap = get_val(row, "chapter", "kategorie", "abschnitt")
            
            if not c_id and not c_title:
                continue

            controls.append({
                "control_id": c_id or "UNKNOWN",
                "title": c_title or "Untitled Control",
                "chapter": c_chap or "General Controls",
                "type": get_val(row, "type", "art") or "Preventive",
                "concept": get_val(row, "concept", "cybersecurity_concept") or "Protect",
                "capability": get_val(row, "capability") or "Information protection",
                "domain": get_val(row, "domain") or "Governance",
                "requirement": c_req or "No explicit requirement provided.",
                "sub_controls": get_val(row, "sub_control", "checklist"),
                "evidence": get_val(row, "evidence", "nachweis"),
                "common_gaps": get_val(row, "common_gap", "schwachstelle"),
                "consult_roles": get_val(row, "consult", "rollen", "responsible"),
                "compensating_controls": get_val(row, "compensating", "ersatzmassnahme")
            })
            
        return controls

    @staticmethod
    def export_annotated_gap_report(
        output_path: str,
        project_title: str,
        gap_assessments: List[Dict[str, Any]]
    ) -> str:
        """
        Exports an executive, audit-ready Excel workbook highlighting assessed controls,
        identified gaps in plain English, agreed action plans, and selected compensating controls.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ISO 27002 Gap Assessment"

        header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid") # Civic Cobalt
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        mono_font = Font(name="Consolas", size=10, bold=True)
        data_font = Font(name="Segoe UI", size=10)
        border_side = Side(border_style="thin", color="D8D6CE")
        cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        wrap_alignment = Alignment(vertical="top", wrap_text=True)

        # Status Fills
        status_fills = {
            "COMPLIANT": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"), # Light Emerald
            "PARTIAL_GAP": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"), # Light Amber
            "CRITICAL_GAP": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"), # Light Crimson
            "COMPENSATING_CONTROL": PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid") # Light Indigo
        }

        columns = [
            ("Chapter", 20),
            ("Control ID", 14),
            ("Control Title", 28),
            ("Compliance Status", 22),
            ("Policy vs Reality Drift", 35),
            ("Plain English Gap Breakdown", 45),
            ("Actionable Next Steps / To-Dos", 40),
            ("Consultation Roles", 30),
            ("Risk Compromise / Compensating Control", 40),
            ("User Decision & Notes", 35)
        ]

        for col_idx, (col_name, width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        ws.row_dimensions[1].height = 28

        for row_idx, item in enumerate(gap_assessments, start=2):
            ws.cell(row=row_idx, column=1, value=item.get("chapter", "")).font = data_font
            
            id_cell = ws.cell(row=row_idx, column=2, value=item.get("control_id", ""))
            id_cell.font = mono_font
            id_cell.alignment = Alignment(horizontal="center", vertical="top")

            ws.cell(row=row_idx, column=3, value=item.get("control_title", "")).font = Font(name="Segoe UI", size=10, bold=True)
            
            status = item.get("status", "PARTIAL_GAP")
            status_cell = ws.cell(row=row_idx, column=4, value=status.replace("_", " "))
            status_cell.font = Font(name="Segoe UI", size=10, bold=True)
            status_cell.alignment = Alignment(horizontal="center", vertical="top")
            if status in status_fills:
                status_cell.fill = status_fills[status]

            reality_summary = f"Policy: {item.get('policy_claim', 'None documented')}\nActual: {item.get('operational_reality', 'Not observed')}"
            ws.cell(row=row_idx, column=5, value=reality_summary).font = data_font
            ws.cell(row=row_idx, column=6, value=item.get("plain_english_gap", "")).font = data_font

            todos = item.get("actionable_todos", [])
            todos_str = "\n".join([f"• {t}" for t in todos]) if isinstance(todos, list) else str(todos)
            ws.cell(row=row_idx, column=7, value=todos_str).font = data_font

            roles = item.get("roles_to_consult", [])
            roles_str = ", ".join(roles) if isinstance(roles, list) else str(roles)
            ws.cell(row=row_idx, column=8, value=roles_str).font = data_font

            ws.cell(row=row_idx, column=9, value=str(item.get("compensating_controls", ""))).font = data_font
            ws.cell(row=row_idx, column=10, value=str(item.get("user_decision_notes", ""))).font = data_font

            for col_idx in range(1, 11):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = cell_border
                if col_idx not in [2, 4]:
                    c.alignment = wrap_alignment

            ws.row_dimensions[row_idx].height = 65

        ws.freeze_panes = "A2"
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)
        return output_path
